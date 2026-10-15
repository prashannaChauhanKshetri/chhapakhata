"""Read go-live data from CSV or Excel files with a header row.

Templates live in docs/templates/. Column names are matched case-insensitively
and spaces become underscores, so "Opening Balance" and "opening_balance" both work.
"""

from __future__ import annotations

import csv
from collections.abc import Iterator
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _key(name: Any) -> str:
    return str(name or "").strip().lower().replace(" ", "_").replace("/", "_")


def read_rows(path: str | Path) -> Iterator[tuple[int, dict[str, Any]]]:
    """Yield ``(row_number, {column: value})`` for every non-empty data row."""
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [_key(c) for c in next(rows)]
        for number, values in enumerate(rows, start=2):
            row = {h: v for h, v in zip(header, values, strict=False) if h}
            if any(v not in (None, "") for v in row.values()):
                yield number, row
        return
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        for number, raw in enumerate(reader, start=2):
            row = {_key(k): (v.strip() if isinstance(v, str) else v) for k, v in raw.items()}
            if any(v not in (None, "") for v in row.values()):
                yield number, row


def text(row: dict[str, Any], key: str) -> str:
    value = row.get(key)
    return "" if value is None else str(value).strip()


def decimal(row: dict[str, Any], key: str, default: Decimal | None = None) -> Decimal | None:
    value = row.get(key)
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"'{key}' must be a number, got '{value}'") from exc


def boolean(row: dict[str, Any], key: str, default: bool = True) -> bool:
    value = text(row, key).lower()
    if not value:
        return default
    return value in {"1", "yes", "y", "true", "active"}
