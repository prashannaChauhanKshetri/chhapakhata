"""Backups and data export.

Backups: ``pg_dump`` custom-format dump plus a tarball of uploaded media,
with a SHA-256 checksum. Retention keeps 14 daily, 8 weekly and 12 monthly
sets. If ``RCLONE_REMOTE`` is configured (ideally an rclone *crypt* remote),
the backup folder is copied online. ``rclone copy`` never deletes remote
files, so a wiped local disk can never wipe the online copy.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import logging
import os
import re
import shutil
import subprocess
import tarfile
from pathlib import Path
from typing import Any

from django.apps import apps
from django.conf import settings
from django.db import connection
from django.utils import timezone
from openpyxl import Workbook

from apps.core.activity import log_activity
from apps.tenancy.context import current_business_id

from .models import BackupRun, BackupStatus, BackupTrigger, OffsiteStatus

logger = logging.getLogger(__name__)

STAMP_RE = re.compile(r"chhapakhata-(db|media)-(\d{8}-\d{6})\.")
KEEP_DAILY, KEEP_WEEKLY, KEEP_MONTHLY = 14, 8, 12


def _bin(name: str) -> str:
    folder = os.environ.get("PG_BIN_DIR", "")
    candidate = Path(folder) / name if folder else None
    if candidate and candidate.exists():
        return str(candidate)
    found = shutil.which(name)
    if not found:
        raise RuntimeError(f"'{name}' was not found. Install the PostgreSQL client tools.")
    return found


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _dump_database(target: Path) -> None:
    db = connection.settings_dict
    # Every business's rows: the app role is under row-level security, so the
    # dump runs with the policies on and the platform bypass set.
    env = {
        **os.environ,
        "PGPASSWORD": str(db.get("PASSWORD") or ""),
        "PGOPTIONS": "-c app.rls_bypass=on",
    }
    cmd = [
        _bin("pg_dump"),
        "--format=custom",
        "--no-owner",
        "--enable-row-security",
        "--file",
        str(target),
        "--dbname",
        db["NAME"],
    ]
    if db.get("HOST"):
        cmd += ["--host", str(db["HOST"])]
    if db.get("PORT"):
        cmd += ["--port", str(db["PORT"])]
    if db.get("USER"):
        cmd += ["--username", str(db["USER"])]
    subprocess.run(cmd, check=True, capture_output=True, timeout=1800, env=env)  # noqa: S603


def _archive_media(target: Path) -> bool:
    media = Path(settings.MEDIA_ROOT)
    if not media.exists() or not any(media.iterdir()):
        return False
    with tarfile.open(target, "w:gz") as tar:
        tar.add(media, arcname="media")
    return True


def prune_backups(folder: Path, today: dt.date | None = None) -> list[Path]:
    """Delete backup sets outside the daily/weekly/monthly retention windows."""
    today = today or timezone.localdate()
    sets: dict[str, list[Path]] = {}
    for path in folder.glob("chhapakhata-*"):
        match = STAMP_RE.match(path.name)
        if match:
            sets.setdefault(match.group(2), []).append(path)

    keep: set[str] = set()
    stamps = sorted(sets, reverse=True)
    days = {s: dt.datetime.strptime(s, "%Y%m%d-%H%M%S").date() for s in stamps}
    seen_weeks, seen_months = set(), set()
    for stamp in stamps:
        day = days[stamp]
        age = (today - day).days
        if age < KEEP_DAILY:
            keep.add(stamp)
        week = day.isocalendar()[:2]
        if age < KEEP_WEEKLY * 7 and week not in seen_weeks:
            keep.add(stamp)
            seen_weeks.add(week)
        month = (day.year, day.month)
        if age < KEEP_MONTHLY * 31 and month not in seen_months:
            keep.add(stamp)
            seen_months.add(month)

    removed = []
    for stamp, paths in sets.items():
        if stamp in keep:
            continue
        for path in paths:
            path.unlink(missing_ok=True)
            removed.append(path)
    return removed


def sync_offsite(run: BackupRun) -> BackupRun:
    remote = settings.RCLONE_REMOTE
    if not remote:
        run.offsite_status = OffsiteStatus.SKIPPED
        run.offsite_message = "Online copy is not configured."
        run.save(update_fields=["offsite_status", "offsite_message"])
        return run
    rclone = shutil.which("rclone")
    if rclone is None:
        run.offsite_status = OffsiteStatus.FAILED
        run.offsite_message = "rclone is not installed."
        run.save(update_fields=["offsite_status", "offsite_message"])
        return run
    cmd = [
        rclone,
        "copy",
        str(settings.BACKUP_DIR),
        remote,
        "--retries",
        "3",
        "--low-level-retries",
        "5",
    ]
    if settings.RCLONE_CONFIG:
        cmd += ["--config", settings.RCLONE_CONFIG]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=3600)  # noqa: S603
        run.offsite_status = OffsiteStatus.SUCCESS
        run.offsite_message = f"Copied to {remote}"
    except subprocess.CalledProcessError as exc:
        stderr = (exc.stderr or b"").decode(errors="replace")
        offline = "no such host" in stderr or "dial tcp" in stderr or "network" in stderr
        run.offsite_status = OffsiteStatus.PENDING if offline else OffsiteStatus.FAILED
        run.offsite_message = (
            "No internet; will retry with the next backup." if offline else stderr[-250:]
        )
    except subprocess.TimeoutExpired:
        run.offsite_status = OffsiteStatus.PENDING
        run.offsite_message = "Upload timed out; will retry with the next backup."
    run.save(update_fields=["offsite_status", "offsite_message"])
    return run


def run_backup(trigger: str = BackupTrigger.SCHEDULED, user: Any = None) -> BackupRun:
    """Create one backup set, prune old sets and copy online when possible."""
    run = BackupRun.objects.create(trigger=trigger, created_by=user)
    folder = Path(settings.BACKUP_DIR) / "sets"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M%S")
    db_file = folder / f"chhapakhata-db-{stamp}.dump"
    media_file = folder / f"chhapakhata-media-{stamp}.tar.gz"
    try:
        _dump_database(db_file)
        has_media = _archive_media(media_file)
        checksum = _sha256(db_file)
        (folder / f"chhapakhata-db-{stamp}.sha256").write_text(f"{checksum}  {db_file.name}\n")
        run.database_file = str(db_file)
        run.media_file = str(media_file) if has_media else ""
        run.size_bytes = db_file.stat().st_size + (media_file.stat().st_size if has_media else 0)
        run.sha256 = checksum
        run.status = BackupStatus.SUCCESS
        removed = prune_backups(folder)
        run.message = f"Removed {len(removed)} old file(s)." if removed else ""
    except Exception as exc:
        logger.exception("Backup failed")
        detail = (
            exc.stderr.decode(errors="replace")
            if isinstance(exc, subprocess.CalledProcessError) and exc.stderr
            else str(exc)
        )
        run.status = BackupStatus.FAILED
        run.message = detail[-2000:]
    run.finished_at = timezone.now()
    run.save()
    if run.status == BackupStatus.SUCCESS:
        sync_offsite(run)
    if current_business_id() is not None:  # the nightly backup belongs to no one business
        log_activity(user, f"backup.{run.status}", run, message=run.message[:200])
    return run


def scheduled_backup() -> int:
    """Entry point for the django-q schedule."""
    return run_backup(BackupTrigger.SCHEDULED).pk


# ---------------------------------------------------------------------------
# Data export
# ---------------------------------------------------------------------------

EXPORT_MODELS = [
    ("customers", "Customer", ["created_by"]),
    ("catalog", "MediaType", []),
    ("catalog", "AddOn", []),
    ("orders", "Order", []),
    ("orders", "OrderItem", []),
    ("orders", "OrderItemAddOn", []),
    ("production", "Job", []),
    ("production", "JobStageEvent", []),
    ("billing", "Invoice", []),
    ("billing", "InvoiceLine", []),
    ("billing", "Payment", []),
    ("billing", "LedgerEntry", []),
    ("inventory", "StockItem", []),
    ("inventory", "Roll", []),
    ("inventory", "Supplier", []),
    ("inventory", "Purchase", []),
    ("inventory", "PurchaseLine", []),
    ("inventory", "StockMovement", []),
    ("core", "ActivityLog", []),
]


def _cell(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, dict | list):
        return str(value)
    return value


def export_all_data() -> bytes:
    """The current business's tables as one Excel workbook (one sheet per table).

    Row-level security limits each sheet to the current business; users are
    global, so that sheet is filtered here.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for app_label, model_name, exclude in EXPORT_MODELS:
        model = apps.get_model(app_label, model_name)
        fields = [f for f in model._meta.concrete_fields if f.name not in exclude]
        ws = wb.create_sheet(title=model._meta.verbose_name_plural.title()[:31])
        ws.append([f.attname for f in fields])
        for row in model.objects.order_by("pk").values_list(*[f.attname for f in fields]):
            ws.append([_cell(v) for v in row])
    users = apps.get_model("accounts", "User")
    ws = wb.create_sheet(title="Users")
    columns = [
        "id",
        "username",
        "first_name",
        "last_name",
        "role",
        "phone",
        "is_active",
        "last_login",
    ]
    ws.append(columns)
    staff = users.objects.filter(business_id=current_business_id())
    for row in staff.order_by("pk").values_list(*columns):
        ws.append([_cell(v) for v in row])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
