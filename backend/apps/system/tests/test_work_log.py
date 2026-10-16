"""Importing and exporting the shop's Excel work log."""

import datetime as dt
import io
from decimal import Decimal as D

import pytest
from freezegun import freeze_time
from openpyxl import Workbook, load_workbook

from apps.billing.models import LedgerEntryType, Payment
from apps.billing.selectors import customer_balance
from apps.catalog.models import MediaType, PricingUnit
from apps.customers.models import Customer
from apps.orders.models import Order, OrderSource, OrderStatus
from apps.orders.selectors import last_rates_for
from apps.system.work_log import export_work_log, import_work_log, parse_date, parse_mode

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def _today():
    with freeze_time("2026-10-16 10:00"):
        yield


HEADER = [
    "Date",
    "Particular",
    "Address",
    "Items",
    "Height",
    "Width",
    "Qty",
    "SQFT",
    "Rate",
    "Amount",
    "Details",
    "Payment S",
    "Payment M",
    "Discount",
    "Balance",
]
ROWS = [
    [
        "2083.04.01",
        "Aayan Digital",
        "Nepalgunj",
        "Korean Banner",
        2.5,
        4,
        1,
        10,
        50,
        500,
        "Photo",
        500,
        None,
        0,
        0,
    ],
    [
        "2083.04.01",
        "Aayan Digital",
        "Nepalgunj",
        "Normal Photo",
        None,
        None,
        2,
        0,
        160,
        320,
        "Photo",
        0,
        None,
        0,
        320,
    ],
    [
        "2083.04.01",
        "Vianet Communication",
        "Kathmandu",
        "Sticker",
        2,
        3,
        3,
        18,
        80,
        1440,
        "Khajura",
        0,
        None,
        0,
        1440,
    ],
    [
        "2083.04.02",
        "Anil Sir",
        "Nepalgunj",
        "Banner",
        5,
        4,
        1,
        20,
        30,
        600,
        "Name",
        300,
        "eSewa",
        50,
        250,
    ],
]


def _xlsx(rows, header=HEADER, title_rows=0):
    wb = Workbook()
    ws = wb.active
    for _ in range(title_rows):
        ws.append(["OM Traders work details"])
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _run(owner, rows=ROWS, title_rows=0, **kwargs):
    return import_work_log(
        _xlsx(rows, title_rows=title_rows), name="Work_Details_export.xlsx", user=owner, **kwargs
    )


def test_dry_run_shows_everything_and_saves_nothing(owner):
    result = _run(owner)
    assert result.dry_run is True
    assert result.rows_imported == 4
    assert result.orders == 3  # two rows for Aayan Digital on the same day become one order
    assert sorted(result.new_customers) == ["Aayan Digital", "Anil Sir", "Vianet Communication"]
    assert {i["name"] for i in result.new_items} == {
        "Korean Banner",
        "Normal Photo",
        "Sticker",
        "Banner",
    }
    assert result.totals["amount"] == D("2860.00")
    assert result.totals["paid"] == D("800.00")
    assert result.totals["balance"] == D("2010.00")
    assert Order.objects.count() == 0
    assert Customer.objects.count() == 0


def test_import_creates_orders_ledger_and_payments(owner):
    result = _run(owner, dry_run=False)
    assert result.dry_run is False
    aayan = Customer.objects.get(name="Aayan Digital")
    order = aayan.orders.get()
    assert order.source == OrderSource.IMPORT
    assert order.status == OrderStatus.DELIVERED
    assert order.total == D("820.00")
    assert customer_balance(aayan) == D("320.00")
    assert order.ledger_entries.get().entry_type == LedgerEntryType.IMPORTED_WORK
    anil = Customer.objects.get(name="Anil Sir")
    payment = Payment.objects.get(customer=anil)
    assert payment.mode == "wallet" and payment.wallet == "esewa"
    assert customer_balance(anil) == D("250.00")
    photo = MediaType.objects.get(name="Normal Photo")
    assert photo.pricing_unit == PricingUnit.PIECE
    assert MediaType.objects.get(name="Sticker").pricing_unit == PricingUnit.SQFT
    assert last_rates_for(aayan)[photo.pk].rate == D("160.00")


def test_reimport_only_adds_new_rows(owner):
    _run(owner, dry_run=False)
    extra = [
        "2083.04.04",
        "GM College",
        "Nepalgunj",
        "Banner",
        7,
        5,
        1,
        35,
        60,
        2100,
        "Admission",
        0,
        None,
        0,
        2100,
    ]
    result = _run(owner, rows=[*ROWS, extra], dry_run=False)
    assert result.rows_skipped_duplicate == 4
    assert result.rows_imported == 1
    assert Order.objects.count() == 4
    assert result.matched_customers == 0
    assert Customer.objects.count() == 4


def test_errors_block_the_import_unless_skipped(owner):
    bad = ["2083.13.40", "Broken Row", "", "Banner", 1, 1, 1, 1, 30, 30, "", 0, None, 0, 30]
    result = _run(owner, rows=[*ROWS, bad], dry_run=False)
    assert result.dry_run is True  # nothing saved
    assert result.errors[0].row == 6
    assert Order.objects.count() == 0
    result = _run(owner, rows=[*ROWS, bad], dry_run=False, skip_errors=True)
    assert result.rows_imported == 4
    assert Order.objects.count() == 3


def test_header_can_sit_below_a_title(owner):
    assert _run(owner, title_rows=2).rows_imported == 4


def test_balance_mismatch_is_a_warning(owner):
    row = list(ROWS[0])
    row[-1] = 99
    result = _run(owner, rows=[row])
    assert "Balance" in result.warnings[0].message


def test_parsers():
    assert parse_date("2083.04.01").isoformat() == "2026-07-17"
    assert parse_mode("Khalti") == ("wallet", "khalti")
    assert parse_mode("") == ("cash", "")
    with pytest.raises(ValueError):
        parse_mode("barter")


def test_export_round_trip(owner, make_order):
    _run(owner, dry_run=False)
    make_order(confirm=True)  # an order entered in the system
    content = export_work_log(None, dt.date(2026, 10, 16))
    ws = load_workbook(io.BytesIO(content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Date", "Particular", "Address", "Items")
    assert rows[1][0] == "2083.04.01"
    assert sum(r[14] for r in rows[1:]) > 0
    # re-importing the export adds nothing: imported rows match, system rows carry an order no.
    result = import_work_log(io.BytesIO(content), name="export.xlsx", user=owner, dry_run=True)
    assert result.rows_imported == 0
    assert result.rows_already_in_system == len(rows) - 1


def test_api_permissions_and_upload(api, owner, desk):
    upload = _xlsx(ROWS)
    upload.name = "log.xlsx"
    response = api(owner).post(
        "/api/v1/system/work-log/import/", {"file": upload, "dry_run": "true"}, format="multipart"
    )
    assert response.status_code == 200, response.json()
    assert response.json()["rows_imported"] == 4
    assert (
        api(desk).post("/api/v1/system/work-log/import/", {}, format="multipart").status_code == 403
    )
    assert api(owner).get("/api/v1/system/work-log/export/").status_code == 200
    assert api(owner).get("/api/v1/system/work-log/template/").status_code == 200
    assert api(desk).get("/api/v1/system/work-log/export/").status_code == 403
