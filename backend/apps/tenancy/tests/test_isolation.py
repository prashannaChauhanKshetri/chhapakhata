"""Two businesses on one site never see or touch each other's data.

Shop A is the default business with the usual fixtures (customer Hotel
Kalpana, Star Flex, Test Supplier, an invoiced and part-paid order). Shop B is
a second business with its own owner and a little data of its own. Every
check is made as B's owner against A's records.
"""

from __future__ import annotations

import io
from decimal import Decimal as D

import pytest
from django.apps import apps
from django.db import DatabaseError, connection, transaction
from openpyxl import load_workbook

from apps.billing import services as billing
from apps.core.permissions import Role
from apps.customers.models import Customer
from apps.orders.models import Order
from apps.tenancy.context import business_scope, use_business
from apps.tenancy.models import Business
from apps.tenancy.operations import POLICY
from conftest import make_user

pytestmark = pytest.mark.django_db

# Text that exists only in shop A's data.
A_MARKERS = ["Hotel Kalpana", "9858020005", "Test Supplier", "Star Flex", "Eyelets"]

LIST_URLS = [
    "/api/v1/customers/",
    "/api/v1/media-types/",
    "/api/v1/addons/",
    "/api/v1/orders/",
    "/api/v1/jobs/",
    "/api/v1/jobs/board/",
    "/api/v1/stock-items/",
    "/api/v1/rolls/",
    "/api/v1/suppliers/",
    "/api/v1/purchases/",
    "/api/v1/stock-movements/",
    "/api/v1/invoices/",
    "/api/v1/payments/",
    "/api/v1/ledger/",
    "/api/v1/ledger/summary/",
    "/api/v1/dashboard/",
    "/api/v1/insights/",
    "/api/v1/reports/daily-sales/",
    "/api/v1/reports/monthly-sales/",
    "/api/v1/reports/pending-dues/",
    "/api/v1/reports/collections/",
    "/api/v1/reports/material-vs-sales/",
    "/api/v1/reports/operator-log/",
    "/api/v1/reports/stock/",
    "/api/v1/activity/",
    "/api/v1/users/",
    "/api/v1/users/operators/",
    "/api/v1/settings/shop/",
    "/api/v1/auth/me/",
]


@pytest.fixture
def shop_a(business, owner, desk, operator, star_flex_rolls, make_order, eyelets):
    order = make_order(items=[(D("10"), D("4"), 2)])
    invoice = billing.issue_invoice(order=order, user=desk)
    payment = billing.receive_payment(
        user=desk, customer=order.customer, amount="500", mode="cash", order=order
    )
    return {"business": business, "order": order, "invoice": invoice, "payment": payment}


@pytest.fixture
def shop_b(db):
    shop = Business.objects.create(name="Sharma Prints", slug="sharma-prints")
    with use_business(shop):
        owner_b = make_user("ram", Role.OWNER)
        customer_b = Customer.objects.create(name="Bageshwori Traders", phone="9800000001")
    return {"business": shop, "owner": owner_b, "customer": customer_b}


def _leaks(content: bytes) -> list[str]:
    text = content.decode(errors="replace")
    return [marker for marker in A_MARKERS if marker in text]


def _xlsx_text(content: bytes) -> bytes:
    book = load_workbook(io.BytesIO(content), read_only=True)
    cells = [
        str(value)
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    ]
    return "\n".join(cells).encode()


# -- The database itself --------------------------------------------------------


def test_app_role_is_subject_to_row_level_security():
    with connection.cursor() as cursor:
        cursor.execute("SELECT rolsuper, rolbypassrls FROM pg_roles WHERE rolname = current_user")
        superuser, bypass = cursor.fetchone()
    assert not superuser and not bypass, (
        "Tests must connect as the non-superuser app role (see scripts/db-app-role.sh); "
        "superusers skip row-level security, so isolation would go untested."
    )


def test_every_business_table_has_a_forced_policy():
    tables = {
        model._meta.db_table
        for model in apps.get_models()
        if model._meta.label != "accounts.User"
        and any(
            getattr(f, "related_model", None) is Business and f.name == "business"
            for f in model._meta.concrete_fields
        )
    }
    assert len(tables) >= 26
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT c.relname FROM pg_class c JOIN pg_policies p ON p.tablename = c.relname "
            "WHERE c.relrowsecurity AND c.relforcerowsecurity AND p.policyname = %s",
            [POLICY],
        )
        protected = {row[0] for row in cursor.fetchall()}
    assert tables - protected == set()


def test_without_a_business_tables_look_empty_and_reject_writes(shop_a):
    assert Order.objects.count() == 1
    with business_scope(None):
        assert Order.objects.count() == 0
        with connection.cursor() as cursor:
            cursor.execute("SELECT count(*) FROM billing_ledgerentry")
            assert cursor.fetchone()[0] == 0
        with pytest.raises(DatabaseError), transaction.atomic():
            Customer.objects.create(name="Nobody's customer")


def test_a_row_for_another_business_is_rejected(shop_a, shop_b):
    with pytest.raises(DatabaseError, match="row-level security"), transaction.atomic():
        Customer.objects.create(name="Smuggled", business=shop_b["business"])
    with pytest.raises(DatabaseError, match="row-level security"), transaction.atomic():
        Customer.objects.filter(pk=shop_a["order"].customer_id).update(business=shop_b["business"])


def test_each_business_sees_only_its_rows(shop_a, shop_b):
    with use_business(shop_b["business"]):
        assert list(Customer.objects.values_list("name", flat=True)) == ["Bageshwori Traders"]
        assert Order.objects.count() == 0
    assert "Bageshwori Traders" not in Customer.objects.values_list("name", flat=True)


# -- The API, as shop B's owner ---------------------------------------------------


@pytest.mark.parametrize("url", LIST_URLS)
def test_lists_and_reports_show_nothing_of_another_business(api, shop_a, shop_b, url):
    response = api(shop_b["owner"]).get(url)
    assert response.status_code == 200, (url, response.content[:300])
    assert _leaks(response.content) == []


def test_exports_hold_only_the_business_own_data(api, shop_a, shop_b):
    client = api(shop_b["owner"])
    for url in ["/api/v1/system/export/", "/api/v1/system/work-log/export/"]:
        response = client.get(url)
        assert response.status_code == 200, url
        text = _xlsx_text(response.content)
        assert _leaks(text) == [], url
    users_sheet = load_workbook(io.BytesIO(client.get("/api/v1/system/export/").content))["Users"]
    usernames = [row[1] for row in users_sheet.iter_rows(min_row=2, values_only=True)]
    assert usernames == ["ram"]


def test_another_business_records_are_not_found(api, shop_a, shop_b, owner):
    order, invoice, payment = shop_a["order"], shop_a["invoice"], shop_a["payment"]
    roll = order.items.first().media_type.material.rolls.first()
    urls = [
        f"/api/v1/customers/{order.customer_id}/",
        f"/api/v1/customers/{order.customer_id}/statement/",
        f"/api/v1/customers/{order.customer_id}/rates/",
        f"/api/v1/orders/{order.pk}/",
        f"/api/v1/orders/{order.pk}/quotation-pdf/",
        f"/api/v1/jobs/{order.jobs.first().pk}/",
        f"/api/v1/invoices/{invoice.pk}/",
        f"/api/v1/payments/{payment.pk}/",
        f"/api/v1/payments/{payment.pk}/receipt/",
        f"/api/v1/media-types/{order.items.first().media_type_id}/",
        f"/api/v1/stock-items/{roll.stock_item_id}/",
        f"/api/v1/rolls/{roll.pk}/",
        f"/api/v1/suppliers/{roll.purchase_line.purchase.supplier_id}/",
        f"/api/v1/purchases/{roll.purchase_line.purchase_id}/",
        f"/api/v1/users/{owner.pk}/",
    ]
    client = api(shop_b["owner"])
    for url in urls:
        assert client.get(url).status_code == 404, url
    patch = client.patch(f"/api/v1/customers/{order.customer_id}/", {"name": "x"}, format="json")
    assert patch.status_code == 404
    cancel = client.post(f"/api/v1/invoices/{invoice.pk}/cancel/", {"reason": "Not mine"})
    assert cancel.status_code == 404
    printed = client.post(f"/api/v1/invoices/{invoice.pk}/print/", {"size": "a4"})
    assert printed.status_code == 404


def test_writes_cannot_point_at_another_business_records(api, shop_a, shop_b, operator):
    order = shop_a["order"]
    client = api(shop_b["owner"])
    item = {"media_type": order.items.first().media_type_id, "width": "3", "height": "2"}
    body = {"customer": order.customer_id, "items": [{**item, "quantity": 1}]}
    assert client.post("/api/v1/orders/", body, format="json").status_code == 400

    pay = {"customer": order.customer_id, "order": order.pk, "amount": "100", "mode": "cash"}
    assert client.post("/api/v1/payments/", pay, format="json").status_code == 400

    # B's own order, then try to give its job to A's operator.
    media = client.post("/api/v1/media-types/", {"name": "Flex B", "rate": "30"}, format="json")
    own = {
        "customer": shop_b["customer"].pk,
        "items": [{"media_type": media.json()["id"], "width": "3", "height": "2", "quantity": 1}],
        "confirm": True,
    }
    created = client.post("/api/v1/orders/", own, format="json")
    assert created.status_code == 201, created.content
    with use_business(shop_b["business"]):
        job = Order.objects.get(pk=created.json()["id"]).jobs.get()
    assign = client.patch(f"/api/v1/jobs/{job.pk}/", {"operator": operator.pk}, format="json")
    assert assign.status_code == 400


def test_numbering_and_names_are_per_business(api, shop_a, shop_b):
    client = api(shop_b["owner"])
    # The same rate-card name is fine in another business, but not twice in one.
    first = client.post("/api/v1/media-types/", {"name": "Star Flex", "rate": "30"}, format="json")
    assert first.status_code == 201, first.content
    again = client.post("/api/v1/media-types/", {"name": "Star Flex", "rate": "31"}, format="json")
    assert again.status_code == 400 and "name" in again.json()["errors"][0]["attr"]

    body = {
        "customer": shop_b["customer"].pk,
        "items": [{"media_type": first.json()["id"], "width": "3", "height": "2", "quantity": 1}],
        "confirm": True,
    }
    order_b = client.post("/api/v1/orders/", body, format="json").json()
    assert order_b["order_no"] == shop_a["order"].order_no  # both start at 1
    invoice_b = client.post(f"/api/v1/orders/{order_b['id']}/invoice/", {}, format="json")
    assert invoice_b.json()["invoice_no"] == shop_a["invoice"].invoice_no


def test_users_and_the_user_cap_are_per_business(api, settings, shop_a, shop_b):
    settings.MAX_ACTIVE_USERS = 2  # shop A already has three active users
    client = api(shop_b["owner"])
    users = client.get("/api/v1/users/").json()
    assert [u["username"] for u in users["results"]] == ["ram"]
    body = {"username": "sita", "role": Role.FRONT_DESK, "password": "Long-enough-1"}
    assert client.post("/api/v1/users/", body, format="json").status_code == 201
    third = {**body, "username": "gita"}
    assert client.post("/api/v1/users/", third, format="json").status_code == 400
    with use_business(shop_b["business"]):
        assert shop_b["business"].users.filter(username="sita").exists()


def test_me_and_shop_settings_belong_to_the_business(api, shop_a, shop_b):
    client = api(shop_b["owner"])
    assert client.get("/api/v1/auth/me/").json()["business"] == {
        "id": shop_b["business"].pk,
        "name": "Sharma Prints",
    }
    assert client.get("/api/v1/settings/shop/").json()["name"] == "Sharma Prints"


def test_sign_in_scopes_the_rest_of_the_request(api, shop_b):
    response = api().post("/api/v1/auth/login/", {"username": "ram", "password": "Test@12345"})
    assert response.status_code == 200, response.content
    assert response.json()["business"]["name"] == "Sharma Prints"
    with use_business(shop_b["business"]):
        from apps.core.models import ActivityLog

        assert ActivityLog.objects.filter(action="auth.login").count() == 1


def test_staff_of_a_paused_business_cannot_sign_in(api, shop_b):
    Business.objects.filter(pk=shop_b["business"].pk).update(is_active=False)
    response = api().post("/api/v1/auth/login/", {"username": "ram", "password": "Test@12345"})
    assert response.status_code == 400
    assert "paused" in response.content.decode()


def test_an_account_without_a_business_sees_no_shop_data(api, shop_a):
    admin = make_user("platform", Role.OWNER, is_superuser=True, is_staff=True)
    admin.business = None
    admin.save()
    response = api(admin).get("/api/v1/orders/")
    assert response.status_code == 200 and response.json()["count"] == 0
    assert api(admin).get("/api/v1/settings/shop/").status_code == 403
