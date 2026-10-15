"""Excel and PDF output for :class:`~apps.reports.tables.Report`."""

from __future__ import annotations

import datetime as dt
import io
import re
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.documents import render_pdf
from apps.core.models import ShopSettings
from apps.core.money import format_npr

from .tables import Column, Report

HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
TOTAL_FILL = PatternFill("solid", fgColor="F8FAFC")
THIN = Side(style="thin", color="CBD5E1")
NUMBER_FORMATS = {
    "money": "#,##0.00",
    "sqft": "#,##0.00",
    "number": "#,##0.##",
    "percent": '0.00"%"',
    "date": "yyyy-mm-dd",
}


def _cell_value(value: Any, column: Column) -> Any:
    if isinstance(value, Decimal):
        return float(value) if column.numeric else str(value)
    if isinstance(value, str) and column.numeric:
        try:
            return float(value)
        except ValueError:
            return value
    if isinstance(value, dt.date) and not isinstance(value, dt.datetime):
        return value
    return value


def report_to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[:\\/?*\[\]]", " ", report.title)[:31].strip() or "Report"
    shop = ShopSettings.load()

    ws.append([shop.name])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([report.title])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([report.period])
    ws.append([])

    for metric in report.summary:
        value = metric.value
        ws.append([metric.label, float(value) if isinstance(value, Decimal) else value])
        cell = ws.cell(row=ws.max_row, column=2)
        if metric.kind in NUMBER_FORMATS:
            cell.number_format = NUMBER_FORMATS[metric.kind]
    if report.summary:
        ws.append([])

    widths: dict[int, int] = {}
    for table in report.tables:
        ws.append([table.title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11)
        ws.append([c.label for c in table.columns])
        header_row = ws.max_row
        for idx, column in enumerate(table.columns, start=1):
            cell = ws.cell(row=header_row, column=idx)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(horizontal="right" if column.numeric else "left")
            widths[idx] = max(widths.get(idx, 10), len(column.label) + 2)
        if not table.rows:
            ws.append([table.empty_message])
        for row in table.rows:
            ws.append([_cell_value(row.get(c.key), c) for c in table.columns])
            for idx, column in enumerate(table.columns, start=1):
                cell = ws.cell(row=ws.max_row, column=idx)
                if column.kind in NUMBER_FORMATS:
                    cell.number_format = NUMBER_FORMATS[column.kind]
                widths[idx] = min(
                    48, max(widths.get(idx, 10), len(str(row.get(column.key) or "")) + 2)
                )
        if table.totals:
            ws.append([_cell_value(table.totals.get(c.key), c) for c in table.columns])
            for idx, column in enumerate(table.columns, start=1):
                cell = ws.cell(row=ws.max_row, column=idx)
                cell.font = Font(bold=True)
                cell.fill = TOTAL_FILL
                cell.border = Border(top=THIN)
                if column.kind in NUMBER_FORMATS:
                    cell.number_format = NUMBER_FORMATS[column.kind]
        ws.append([])

    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 12)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 24)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _display(value: Any, kind: str) -> str:
    if value is None or value == "":
        return ""
    if kind == "money":
        return format_npr(value).replace("Rs. ", "")
    if kind in {"sqft", "number"}:
        return f"{Decimal(str(value)):,.2f}".rstrip("0").rstrip(".")
    if kind == "percent":
        return f"{Decimal(str(value)):.2f}%"
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def report_to_pdf(report: Report) -> bytes:
    tables = []
    for table in report.tables:
        tables.append(
            {
                "title": table.title,
                "columns": table.columns,
                "rows": [
                    [
                        {"text": _display(row.get(c.key), c.kind), "numeric": c.numeric}
                        for c in table.columns
                    ]
                    for row in table.rows
                ],
                "totals": (
                    [
                        {"text": _display(table.totals.get(c.key), c.kind), "numeric": c.numeric}
                        for c in table.columns
                    ]
                    if table.totals
                    else None
                ),
                "empty_message": table.empty_message,
            }
        )
    summary = [{"label": m.label, "value": _display(m.value, m.kind)} for m in report.summary]
    landscape = any(len(t.columns) > 6 for t in report.tables)
    return render_pdf(
        "documents/report.html",
        {"report": report, "tables": tables, "summary": summary, "landscape": landscape},
    )
